#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""校验 .md 与 .xlsx 在「9/30 跟进信」涉及字段上是否逐字段一致。"""
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(Path(__file__).resolve().parents[1])
MD = ROOT / "开发信项目/工作区/开发信跟进表.md"
XLSX = ROOT / "开发信项目/工作区/开发信跟进表.xlsx"
# 本轮（9/30 跟进信）改过的行 + 新增「已退信」改标的行
TARGETS = {17, 18, 19, 20, 128, 134, 137, 140, 143, 144, 146, 126, 135, 213,
           13, 63, 82, 90, 104, 106, 151}

md = MD.read_text(encoding="utf-8")
wb = openpyxl.load_workbook(XLSX)


def md_field(num, field):
    m = re.search(rf"^### {num}\.\s.*?(?=^### |\Z)", md, re.S | re.M)
    if not m:
        return None
    # ⚠️ 字段名可能带**对齐补空格**（`| 状态     |`）—— 2026-10-02 起全表改成这种格式，
    #    所以必须用 `\s*` 容忍，不能用 `^| 字段 |`。
    pat = re.compile(rf"^\|\s*{field}\s*\|\s*(.*?)\s*\|\s*$", re.M)
    hit = pat.search(m.group(0))
    if not hit:
        return None
    return hit.group(1).replace("**", "").split("（")[0].strip()


bad = 0
checked = 0
for ws in wb.worksheets:
    for r in range(5, ws.max_row + 1):
        n = ws.cell(r, 1).value
        if n not in TARGETS:
            continue
        for col, f in ((13, "状态"), (14, "跟进次数"), (15, "下次跟进日"), (16, "最后联系日")):
            x = str(ws.cell(r, col).value)
            m = md_field(n, f)
            checked += 1
            # ⚠️ 既有表示差异（非本次引入）：`.xlsx` 的「跟进次数」用数字 **0** 表示「未跟进过」
            #    （全表 154 行都是 0），而 `.md` 用 `—`。两者语义相同，视为一致。
            if f == "跟进次数" and x == "0" and m == "—":
                continue
            if x != m:
                print(f"  ❌ #{n} [{ws.title}] {f}: xlsx={x!r}  md={m!r}")
                bad += 1

print(f"检查 {checked} 个字段，差异 {bad} 处")
print("✅ md ↔ xlsx 逐字段一致" if bad == 0 else "⚠️ 存在差异")
sys.exit(1 if bad else 0)
