#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""阶段 3：校验 batch4 补登记结果"""
import re, sys, zipfile, json
from pathlib import Path
import openpyxl

ROOT = Path("/Users/eric/WorkBuddy AI/2026-09-27-00-43-30/jzp-americas-dev")
XLSX = ROOT / "开发信项目/工作区/开发信跟进表.xlsx"
BAK = ROOT / ".workbuddy-ai/开发信跟进表.bak-20260929-补batch4前.xlsx"
OLD, NEW = "开发信跟进表", "开发信跟进表-batch5+6"
ok = True


def chk(label, cond, detail=""):
    global ok
    print(f"  {'✅' if cond else '❌'} {label}{('  ' + detail) if detail else ''}")
    if not cond:
        ok = False


wb = openpyxl.load_workbook(XLSX)
bb = openpyxl.load_workbook(BAK)

print("=" * 70)
print("① 老 sheet 追加 31 行")
wo, bo = wb[OLD], bb[OLD]
chk("max_row = 149（118 + 31）", wo.max_row == 149, f"实际 {wo.max_row}")
nums = [wo.cell(r, 1).value for r in range(119, 150)]
chk("# 列为 117–147", nums == list(range(117, 148)), f"{nums[0]}–{nums[-1]}")
chk("新增行公司名无空", all(wo.cell(r, 3).value for r in range(119, 150)))
# 样式比对：**比渲染结果，不比原始 style 索引**
#   openpyxl 复制样式时可能新建一条 fill 记录（fillId 4 → 40），
#   颜色/边框/数字格式完全一致 —— 这是无害的重复，不该判为失败（实测踩过）
def vis(c):
    f = c.fill
    b = c.border
    a = c.alignment
    return (f.fill_type, str(f.fgColor.rgb),
            b.left.style, b.right.style, b.top.style, b.bottom.style,
            c.number_format, a.horizontal, a.vertical, a.wrap_text,
            c.font.name, c.font.size, c.font.bold, str(c.font.color.rgb) if c.font.color else None)


bad = []
for r in range(119, 150):
    for j in range(1, 18):
        if vis(wo.cell(r, j)) != vis(wo.cell(6, j)):
            bad.append((r, j))
chk("新增行样式（渲染结果）= 第 6 行（527 格）", not bad, f"不一致 {len(bad)} 格 {bad[:3]}")
gray = sum(1 for r in range(119, 150) for j in (14, 15, 16)
           if wo.cell(r, j).fill.start_color.rgb in ("00F5F7F9", "FFF5F7F9", "F5F7F9"))
chk("N/O/P 灰底（应 93 格）", gray == 93, f"实际 {gray}")
dvs = sorted(str(d.sqref) for d in wo.data_validations.dataValidation)
chk("下拉 sqref 覆盖新行（:200）", all(s.endswith("200") for s in dvs), str(dvs))
cfs = [str(x.sqref) for x in wo.conditional_formatting]
chk("条件格式覆盖新行（M5:M200）", cfs == ["M5:M200"], str(cfs))
chk("AutoFilter 已修正 A4:Q149", str(wo.auto_filter.ref) == "A4:Q149", str(wo.auto_filter.ref))

print()
print("② 老 sheet 前 118 行未被改动")
diff = 0
for r in range(1, 119):
    for j in range(1, 18):
        if wo.cell(r, j).value != bo.cell(r, j).value:
            diff += 1
chk("前 118 行单元格值一致", diff == 0, f"差异 {diff} 格")
chk("冻结窗格一致", str(wo.freeze_panes) == str(bo.freeze_panes), str(wo.freeze_panes))
chk("合并单元格一致", sorted(map(str, wo.merged_cells.ranges)) == sorted(map(str, bo.merged_cells.ranges)))

print()
print("③ 新 sheet 重编号")
wn = wb[NEW]
chk("max_row 仍 = 96", wn.max_row == 96, f"实际 {wn.max_row}")
nn = [wn.cell(r, 1).value for r in range(5, 97)]
chk("# 列为 148–239", nn == list(range(148, 240)), f"{nn[0]}–{nn[-1]}")
chk("新 sheet 下拉仍 4 个 :200", all(str(d.sqref).endswith("200") for d in wn.data_validations.dataValidation))

print()
print("④ 完整性")
z = zipfile.ZipFile(XLSX)
chk("testzip 通过", z.testzip() is None)
pa = re.search(r"<pane[^>]*/>", z.read("xl/worksheets/sheet1.xml").decode("utf-8"))
pb = re.search(r"<pane[^>]*/>", zipfile.ZipFile(BAK).read("xl/worksheets/sheet1.xml").decode("utf-8"))
chk("老 sheet <pane> 字节不变", (pa.group(0) if pa else None) == (pb.group(0) if pb else None),
    (pa.group(0) if pa else "无"))

print()
print("⑤ 数据抽查")
rows = json.loads(Path("/tmp/batch4_rows.json").read_text(encoding="utf-8"))
mis = sum(1 for i, r in enumerate(rows) if wo.cell(119 + i, 3).value != r["公司名称"])
chk("31 行与构建数据一致", mis == 0, f"不一致 {mis} 行")
st = {}
for r in rows:
    st[r["状态"]] = st.get(r["状态"], 0) + 1
print(f"     状态分布：{st}")

print()
print("=" * 70)
print("✅ 全部通过" if ok else "⚠️ 有未通过项")
sys.exit(0 if ok else 1)
