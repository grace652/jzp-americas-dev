#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""10/02 跟进批次 · 跟进表更新（.md + .xlsx 双份同步）

31 封已投递成功（守护进程 3h12m 跑完，0 失败，INBOX 无退信）。
  · 跟进次数 +1
  · 最后联系日 = 2026-10-02
  · 下次跟进日 = 首封日 + N 个工作日（次数=1 → +8bd；次数=2 → +15bd）
"""
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(Path(__file__).resolve().parents[1])
XLSX = ROOT / "开发信项目/工作区/开发信跟进表.xlsx"
MD = ROOT / "开发信项目/工作区/开发信跟进表.md"
WORK = ROOT / ".workbuddy-ai/今日工作单-2026-10-02.json"
C_STATUS, C_CNT, C_NEXT, C_LAST, C_NOTE = 13, 14, 15, 16, 17


def add_bd(d, n):
    cur = d
    while n > 0:
        cur += timedelta(days=1)
        if cur.weekday() < 5:
            n -= 1
    return cur


def pdate(s):
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})", str(s or ""))
    return date(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def main():
    import json
    work = json.loads(WORK.read_text(encoding="utf-8"))
    wb = openpyxl.load_workbook(XLSX)
    rows = {}
    for ws in wb.worksheets:
        for r in range(5, ws.max_row + 1):
            if ws.cell(r, 1).value is not None:
                rows[ws.cell(r, 1).value] = (ws, r)

    md = MD.read_text(encoding="utf-8")

    def blk_raw(num):
        m = re.search(rf"^### {num}\.\s.*?(?=^### |\Z)", md, re.S | re.M)
        return m.group(0) if m else ""

    log = []
    for x in work:
        n = x["num"]
        if n not in rows:
            print(f"  ⚠️ #{n} 表内找不到"); continue
        ws, r = rows[n]
        old = int(ws.cell(r, C_CNT).value or 0)
        new = old + 1
        first = pdate(ws.cell(r, 2).value)
        nxt = add_bd(first, 8 if new == 1 else 15) if first else None
        ws.cell(r, C_CNT).value = new
        ws.cell(r, C_LAST).value = "2026-10-02"
        if nxt:
            ws.cell(r, C_NEXT).value = nxt.isoformat()
        b0 = blk_raw(n)
        ft = re.sub(r"[（(].*", "", re.search(r"\*\*信文件\*\*：(.+)", b0).group(1)
                    ).strip().strip("`").replace("clients/", "")
        slug = ft.replace("开发信-", "开发信-跟进-")
        if "第 2 次" in x["round"]:
            slug = slug.replace("开发信-跟进-", "开发信-跟进-第2轮-")
        add = (f"✅ 10/02 跟进信已发（clients/{slug or '开发信-跟进-*.md'}）｜"
               f"守护进程 3h12m 跑完，31/31 成功 0 失败，INBOX 无退信")
        cur = ws.cell(r, C_NOTE).value or ""
        ws.cell(r, C_NOTE).value = (cur + "｜" + add) if cur and cur != "—" else add
        log.append((n, old, new, nxt))
    wb.save(XLSX)
    print(f"✅ xlsx：{len(log)} 家已更新")

    def blk(t, num):
        m = re.search(rf"^### {num}\.\s.*?(?=^### |\Z)", t, re.S | re.M)
        return m.group(0) if m else None

    changed = 0
    for n, old, new, nxt in log:
        ws, r = rows[n]
        b = blk(md, n)
        if b is None:
            continue
        nb = b
        for f, v in (("跟进次数", f"**{new}**"), ("下次跟进日", nxt.isoformat() if nxt else ""),
                     ("最后联系日", "**2026-10-02**")):
            nb = re.sub(rf"^\|\s*{f}\s*\|\s*(.*?)\s*\|\s*$", f"| {f} | {v} |", nb, count=1, flags=re.M)
        note = ws.cell(r, C_NOTE).value or ""
        nb = re.sub(r"^\|\s*备注\s*\|\s*(.*?)\s*\|\s*$",
                    lambda m: "| 备注 | " + note + " |", nb, count=1, flags=re.M)
        if nb != b:
            md = md.replace(b, nb, 1); changed += 1
    MD.write_text(md, encoding="utf-8")
    print(f"✅ md：{changed} 个块已同步")
    for n, old, new, nxt in sorted(log)[:4]:
        print(f"   #{n}: 次数 {old}→{new}  下次 {nxt}")
    print(f"   …（共 {len(log)} 家）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
